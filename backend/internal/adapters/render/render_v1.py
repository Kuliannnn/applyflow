"""Template 1: plain, text-first A4 exports. No HTML, URLs, or remote assets."""
import io
import json
import os
import sys
from datetime import datetime, timezone
from html import escape

import reportlab
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

MAX_BYTES = 10 * 1024 * 1024


class ExportError(Exception):
    pass


def fonts():
    directory = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    for name, filename in [("ExportRegular", "Vera.ttf"), ("ExportBold", "VeraBd.ttf")]:
        pdfmetrics.registerFont(TTFont(name, os.path.join(directory, filename)))


def blocks(content):
    if content["kind"] == "resume":
        yield "title", content["title"]
        for section in content["sections"]:
            yield "heading", section["heading"]
            for item in section["items"]:
                yield "item", item["text"]
    else:
        yield "body", content["salutation"]
        for paragraph in content["paragraphs"]:
            yield "body", paragraph["text"]
        yield "closing", content["closing"]


def check_text(items):
    # Fail explicitly rather than emitting black boxes or substituting unsupported names.
    for style, text in items:
        face = pdfmetrics.getFont("ExportBold" if style in ("title", "heading") else "ExportRegular").face
        for char in text:
            if char in "\n\r\t":
                continue
            if ord(char) < 32 or ord(char) not in face.charToGlyph:
                raise ExportError("export_unsupported_character")


class BoundedCanvas(Canvas):
    def showPage(self):
        if self.getPageNumber() > 100:
            raise ExportError("export_page_limit")
        super().showPage()


def pdf_bytes(items):
    output = io.BytesIO()
    body = ParagraphStyle("Body", fontName="ExportRegular", fontSize=10.5, leading=15,
                          textColor=colors.HexColor("#202020"), alignment=TA_LEFT,
                          spaceAfter=8, allowWidows=0, allowOrphans=0, splitLongWords=1)
    styles = {
        "body": body,
        "closing": ParagraphStyle("Closing", parent=body, spaceBefore=10),
        "item": ParagraphStyle("Item", parent=body, leftIndent=10, firstLineIndent=-10, spaceAfter=7),
        "title": ParagraphStyle("Title", parent=body, fontName="ExportBold", fontSize=20,
                                leading=25, spaceAfter=18, keepWithNext=1),
        "heading": ParagraphStyle("Heading", parent=body, fontName="ExportBold", fontSize=12,
                                  leading=17, spaceBefore=12, spaceAfter=6, keepWithNext=1),
    }
    story = []
    for style, text in items:
        text = escape(text.expandtabs(4)).replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br/>")
        if style == "item":
            text = "&#8226; " + text
        story.append(Paragraph(text, styles[style]))

    def footer(canvas, _document):
        canvas.saveState()
        canvas.setFont("ExportRegular", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(canvas.getPageNumber()))
        canvas.restoreState()

    pdf = SimpleDocTemplate(output, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm, title="", author="",
                            subject="", creator="ApplyFlow template 1", pageCompression=1,
                            invariant=1)
    pdf.build(story, onFirstPage=footer, onLaterPages=footer, canvasmaker=BoundedCanvas)
    return output.getvalue()


def docx_bytes(items):
    doc = Document()
    # The bundled Word template has decorative paragraph borders in its styles.
    # Keep exported application documents monochrome and free of inherited rules.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(20)
    sec.footer_distance = Mm(10)
    for name in ["Normal", "Title", "Heading 1", "List Bullet", "Footer"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(32, 32, 32)
        style.font.size = Pt(10.5)
        style.paragraph_format.line_spacing = Pt(15)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 20, 0, 18), ("Heading 1", 12, 12, 6)]:
        style = doc.styles[name]
        style.font.bold = True
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = Pt(size + 5)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for kind, text in items:
        style = {"title": "Title", "heading": "Heading 1", "item": "List Bullet"}.get(kind, "Normal")
        p = doc.add_paragraph(style=style)
        p.add_run(text.expandtabs(4))  # python-docx escapes XML; text never becomes a hyperlink.
        if kind == "closing":
            p.paragraph_format.space_before = Pt(10)
    footer = sec.footer.paragraphs[0]
    doc.styles["Footer"].font.size = Pt(8)
    doc.styles["Footer"].font.color.rgb = RGBColor(102, 102, 102)
    footer.style = doc.styles["Footer"]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    props = doc.core_properties
    props.author = props.last_modified_by = props.comments = props.title = props.subject = props.keywords = ""
    props.created = props.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


def main():
    request = json.loads(sys.stdin.buffer.read(262145))
    fonts()
    items = list(blocks(request["content"]))
    check_text(items)
    # Shared layout preflight bounds both formats, without executing office macros/renderers.
    pdf = pdf_bytes(items)
    result = pdf if request["format"] == "pdf" else docx_bytes(items)
    if not result or len(result) > MAX_BYTES:
        raise ExportError("export_size_limit")
    sys.stdout.buffer.write(result)


if __name__ == "__main__":
    try:
        main()
    except ExportError as exc:
        sys.stderr.write(str(exc))
        sys.exit(2)
    except Exception:
        # Never emit document text, local paths, or a traceback to the worker's logs.
        sys.stderr.write("export_render_failed")
        sys.exit(2)
